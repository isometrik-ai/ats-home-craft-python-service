"""Parse the client Project_data_bulk.xlsx workbook into import records."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from apps.user_service.app.schemas.enums import (
    CommercialUnitType,
    FacilityLocationType,
    FacilityType,
    Facing,
    GateStatus,
    GateType,
    LiftType,
    MeasurementUnit,
    ParkingUserType,
    ParkingVehicleCategory,
    PlotType,
    PropertyType,
    TowerType,
    UnitConfigKind,
    UnitNumberingPattern,
    UnitStatus,
)

HEADER_ROW = 3
DATA_START_ROW = 5

SHEET_PROJECT = "1. Project"
SHEET_TOWERS = "2. Towers"
SHEET_WINGS = "3. Wings"
SHEET_UNIT_TYPES = "4. Unit types"
SHEET_UNITS = "5. Units and plots"
SHEET_AMENITIES = "6. Amenities"
SHEET_PARKING = "7. Parking zones"
SHEET_GATES_LIFTS = "8. Gates and lifts"
SHEET_MEDIA = "9. Photos and files"
SHEET_PINS = "10. Map pins"

PROPERTY_TYPE_MAP: dict[str, list[PropertyType]] = {
    "residential apartments": [PropertyType.RESIDENTIAL],
    "commercial": [PropertyType.COMMERCIAL],
    "plotted development": [PropertyType.PLOTS],
    "residential and commercial": [PropertyType.RESIDENTIAL, PropertyType.COMMERCIAL],
    "apartments and plots": [PropertyType.RESIDENTIAL, PropertyType.PLOTS],
}

AREA_UNIT_MAP: dict[str, MeasurementUnit] = {
    "square feet": MeasurementUnit.SQ_FT,
    "sq_ft": MeasurementUnit.SQ_FT,
    "sqft": MeasurementUnit.SQ_FT,
    "square metres": MeasurementUnit.SQ_M,
    "square meters": MeasurementUnit.SQ_M,
    "sq_m": MeasurementUnit.SQ_M,
    "square yards": MeasurementUnit.GAJ,
    "gaj": MeasurementUnit.GAJ,
}

TOWER_TYPE_MAP: dict[str, TowerType] = {
    "residential": TowerType.RESIDENTIAL,
    "commercial": TowerType.COMMERCIAL,
    "mixed": TowerType.MIXED,
}

NUMBERING_MAP: dict[str, UnitNumberingPattern] = {
    "floor + unit (101, 102)": UnitNumberingPattern.FLOOR_UNIT,
    "floor + unit": UnitNumberingPattern.FLOOR_UNIT,
    "continuous series (1, 2, 3)": UnitNumberingPattern.SEQUENTIAL,
    "continuous series": UnitNumberingPattern.SEQUENTIAL,
    "we will share a custom list": UnitNumberingPattern.CUSTOM,
}

KIND_MAP: dict[str, UnitConfigKind] = {
    "apartment": UnitConfigKind.APARTMENT,
    "shop or office": UnitConfigKind.COMMERCIAL,
    "plot": UnitConfigKind.PLOT,
}

FACING_MAP: dict[str, Facing] = {
    "east": Facing.EAST,
    "west": Facing.WEST,
    "north": Facing.NORTH,
    "south": Facing.SOUTH,
    "north-east": Facing.NORTH_EAST,
    "north east": Facing.NORTH_EAST,
    "north-west": Facing.NORTH_WEST,
    "north west": Facing.NORTH_WEST,
    "south-east": Facing.SOUTH_EAST,
    "south east": Facing.SOUTH_EAST,
    "south-west": Facing.SOUTH_WEST,
    "south west": Facing.SOUTH_WEST,
}

SHOP_TYPE_MAP: dict[str, CommercialUnitType] = {
    "retail shop": CommercialUnitType.RETAIL_SHOP,
    "office": CommercialUnitType.OFFICE,
    "showroom": CommercialUnitType.ANCHOR_STORE,
    "food outlet": CommercialUnitType.FOOD_COURT,
    "other": CommercialUnitType.OTHER,
}

PLOT_TYPE_MAP: dict[str, PlotType] = {
    "residential plot": PlotType.RESIDENTIAL,
    "villa plot": PlotType.VILLA,
    "commercial plot": PlotType.COMMERCIAL,
    "row house": PlotType.VILLA,
}

UNIT_STATUS_MAP: dict[str, UnitStatus] = {
    "available": UnitStatus.VACANT,
    "vacant": UnitStatus.VACANT,
    "not released": UnitStatus.VACANT,
    "owner": UnitStatus.VACANT,
    "booked": UnitStatus.VACANT,
    "sold": UnitStatus.VACANT,
    "blocked": UnitStatus.BLOCKED,
}

AMENITY_TYPE_MAP: dict[str, FacilityType] = {
    "clubhouse": FacilityType.RECREATION,
    "swimming pool": FacilityType.SPORTS,
    "gym": FacilityType.SPORTS,
    "park or garden": FacilityType.RECREATION,
    "parking": FacilityType.PARKING,
    "security or gate": FacilityType.SERVICES,
    "sports": FacilityType.SPORTS,
    "kids play area": FacilityType.RECREATION,
    "temple or religious": FacilityType.SERVICES,
    "other": FacilityType.SERVICES,
}

AMENITY_PLACE_MAP: dict[str, FacilityLocationType] = {
    "inside a tower": FacilityLocationType.IN_TOWER,
    "open area in the project": FacilityLocationType.OUTDOOR_STANDALONE,
    "its own building": FacilityLocationType.INDOOR_CLUBHOUSE,
}

PARKING_USERS_MAP: dict[str, ParkingUserType] = {
    "residents": ParkingUserType.RESIDENT,
    "visitors": ParkingUserType.VISITORS,
}

VEHICLE_MAP: dict[str, ParkingVehicleCategory] = {
    "two-wheeler": ParkingVehicleCategory.TWO_WHEELER,
    "four-wheeler": ParkingVehicleCategory.FOUR_WHEELER,
    "both": ParkingVehicleCategory.BOTH,
}

GATE_TYPE_MAP: dict[str, GateType] = {
    "entry": GateType.ENTRY,
    "exit": GateType.EXIT,
    "entry and exit": GateType.BOTH,
}

GATE_STATUS_MAP: dict[str, GateStatus] = {
    "in use": GateStatus.ACTIVE,
    "closed": GateStatus.INACTIVE,
    "under construction": GateStatus.INACTIVE,
}

LIFT_TYPE_MAP: dict[str, LiftType] = {
    "passenger": LiftType.PASSENGER,
    "service": LiftType.SERVICE,
    "stretcher": LiftType.SERVICE,
    "fire": LiftType.SERVICE,
}

YES_VALUES = {"yes", "true", "1", "y"}
FLOOR_NAME_LEVELS: dict[str, int] = {
    "ground": 0,
    "ground floor": 0,
    "first": 1,
    "first floor": 1,
    "second": 2,
    "second floor": 2,
    "third": 3,
    "third floor": 3,
}


@dataclass
class WorkbookRowError:
    """One invalid or skipped workbook row."""

    sheet: str
    row_number: int
    code: str
    message: str
    raw: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedProjectWorkbook:
    """Normalized records from a Project_data_bulk workbook."""

    projects: list[dict[str, Any]] = field(default_factory=list)
    towers: list[dict[str, Any]] = field(default_factory=list)
    wings: list[dict[str, Any]] = field(default_factory=list)
    unit_types: list[dict[str, Any]] = field(default_factory=list)
    floors: list[dict[str, Any]] = field(default_factory=list)
    units: list[dict[str, Any]] = field(default_factory=list)
    amenities: list[dict[str, Any]] = field(default_factory=list)
    parking_zones: list[dict[str, Any]] = field(default_factory=list)
    gates: list[dict[str, Any]] = field(default_factory=list)
    lifts: list[dict[str, Any]] = field(default_factory=list)
    skipped: list[WorkbookRowError] = field(default_factory=list)
    errors: list[WorkbookRowError] = field(default_factory=list)


def normalize_project_name(value: object) -> str:
    """Collapse whitespace and case-fold for joining project names across sheets."""
    text = _cell_text(value)
    collapsed = " ".join(text.split())
    return unicodedata.normalize("NFKC", collapsed).casefold()


def project_join_key(value: object) -> str:
    """Join key that treats 'Home Kraft' and 'Homekraft' as the same name."""
    return re.sub(r"[^a-z0-9]+", "", normalize_project_name(value))


def _resolve_project_key(parsed: ParsedProjectWorkbook, value: object) -> str:
    """Map a sheet project name onto the parsed project's join key."""
    key = project_join_key(value)
    known = {item["name_key"] for item in parsed.projects}
    if key in known:
        return key
    contained = [item for item in known if key and (key in item or item in key)]
    if len(contained) == 1:
        return contained[0]
    if len(known) == 1:
        return next(iter(known))
    return key


def _cell_text(value: object) -> str:
    """Normalize a spreadsheet cell to stripped text."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.replace("\xa0", " ").strip()
    return str(value).strip()


def _cell_code(value: object) -> str:
    """Unit/plot numbers from Excel must not keep a trailing .0."""
    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = _cell_text(value)
    if re.fullmatch(r"-?\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def _cell_pin(value: object) -> str:
    """Normalize a PIN code cell, dropping Excel float suffixes."""
    if isinstance(value, bool) or value is None:
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _cell_text(value)


def _cell_bool(value: object) -> bool:
    """Treat common yes/true spreadsheet values as True."""
    return _cell_text(value).casefold() in YES_VALUES


def _lookup(mapping: dict[str, Any], value: object) -> Any | None:
    """Map a spreadsheet label onto an enum or coded value."""
    key = _cell_text(value).casefold()
    if not key:
        return None
    if key in mapping:
        return mapping[key]
    for prefix, mapped in mapping.items():
        if key.startswith(prefix):
            return mapped
    return None


def _read_sheet(
    workbook: Any, sheet_name: str
) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    """Read header row 3 and data rows from a named workbook sheet."""
    if sheet_name not in workbook.sheetnames:
        return [], []
    worksheet = workbook[sheet_name]
    headers = [
        _cell_text(cell.value)
        for cell in next(worksheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW))
    ]
    rows: list[tuple[int, dict[str, Any]]] = []
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=DATA_START_ROW, values_only=True),
        start=DATA_START_ROW,
    ):
        if all(item is None or str(item).strip() == "" for item in values):
            continue
        record = {
            headers[index]: values[index] if index < len(values) else None
            for index in range(len(headers))
            if headers[index]
        }
        rows.append((row_number, record))
    return headers, rows


def parse_floor_level(value: object) -> int | None:
    """Parse basement/ground/first-style floor labels into level numbers."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = _cell_text(value).casefold()
    if text in FLOOR_NAME_LEVELS:
        return FLOOR_NAME_LEVELS[text]
    match = re.search(r"-?\d+", text)
    if match:
        return int(match.group(0))
    return None


def floor_display_name(level: int) -> str:
    """Import-only floor label from Floor number. Ignores workbook Floor name words."""
    if level == 0:
        return "Ground"
    if level < 0:
        return f"Basement {abs(level)}"
    return f"Floor {level}"


def unit_codes_for_floor(
    *,
    starting_unit_number: int,
    floor_level: int,
    units_per_floor: int,
    numbering_pattern: UnitNumberingPattern | str,
) -> list[str]:
    """Build unit codes for one above-ground floor from the Towers sheet numbering."""
    if units_per_floor <= 0 or floor_level <= 0:
        return []
    pattern = (
        numbering_pattern.value
        if isinstance(numbering_pattern, UnitNumberingPattern)
        else str(numbering_pattern)
    )
    if pattern == UnitNumberingPattern.CUSTOM.value:
        return []
    if pattern == UnitNumberingPattern.SEQUENTIAL.value:
        base = starting_unit_number + (floor_level - 1) * units_per_floor
        return [str(base + slot) for slot in range(units_per_floor)]
    remainder = starting_unit_number % 100
    stride = 100 if units_per_floor >= 10 or remainder <= units_per_floor else 10
    base = starting_unit_number + (floor_level - 1) * stride
    return [str(base + slot) for slot in range(units_per_floor)]


def build_tower_floor_specs(
    *,
    basement_count: int,
    upper_floor_count: int,
    parking_levels: set[int] | None = None,
) -> list[dict[str, Any]]:
    """Floor payloads the tower create API already accepts as nested floors."""
    parking_levels = parking_levels or set()
    specs: list[dict[str, Any]] = []
    for level in range(1, int(basement_count or 0) + 1):
        specs.append(
            {
                "level_number": -level,
                "display_name": floor_display_name(-level),
                "is_parking": -level in parking_levels,
            }
        )
    for level in range(1, int(upper_floor_count or 0) + 1):
        specs.append(
            {
                "level_number": level,
                "display_name": floor_display_name(level),
                "is_parking": level in parking_levels,
            }
        )
    return specs


def _parse_possession(value: object) -> date | None:
    """Parse expected-possession dates from common spreadsheet formats."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _cell_text(value)
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%b %Y", "%B %Y"):
        try:
            parsed = datetime.strptime(text, fmt)
            if "%Y" in fmt and "%d" not in fmt:
                return parsed.date().replace(day=1)
            return parsed.date()
        except ValueError:
            continue
    return None


def _parse_float(value: object) -> float | None:
    """Parse a numeric spreadsheet cell, ignoring commas."""
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = _cell_text(value).replace(",", "")
    try:
        return float(text)
    except ValueError:
        return None


def _parse_int(value: object, default: int | None = None) -> int | None:
    """Parse an integer spreadsheet cell, using default when empty."""
    number = _parse_float(value)
    if number is None:
        return default
    return int(number)


def _slug(value: str, fallback: str = "item") -> str:
    """Build a short code from a display name."""
    cleaned = re.sub(r"[^a-z0-9]+", "-", value.casefold()).strip("-")
    return (cleaned or fallback)[:64]


def _infer_bedrooms(type_name: str) -> float:
    """Infer bedroom count from a type name such as 3+Study."""
    match = re.match(r"(\d+)", type_name.strip())
    return float(match.group(1)) if match else 1.0


def parse_project_workbook(payload: bytes) -> ParsedProjectWorkbook:
    """Load an xlsx workbook and apply Happy Trails import rules."""
    workbook = load_workbook(BytesIO(payload), data_only=True, read_only=True)
    parsed = ParsedProjectWorkbook()
    try:
        _parse_projects(workbook, parsed)
        _parse_towers(workbook, parsed)
        _parse_wings(workbook, parsed)
        _parse_unit_types(workbook, parsed)
        _parse_units(workbook, parsed)
        _parse_amenities(workbook, parsed)
        _parse_parking_zones(workbook, parsed)
        _parse_gates_lifts(workbook, parsed)
        _skip_optional_sheet(workbook, parsed, SHEET_MEDIA, "media_deferred")
        _skip_optional_sheet(workbook, parsed, SHEET_PINS, "map_pins_skipped")
    finally:
        workbook.close()
    return parsed


def _parse_projects(workbook: Any, parsed: ParsedProjectWorkbook) -> None:
    """Parse the Project sheet into create_project payloads."""
    _, rows = _read_sheet(workbook, SHEET_PROJECT)
    if not rows:
        parsed.errors.append(
            WorkbookRowError(
                sheet=SHEET_PROJECT,
                row_number=HEADER_ROW,
                code="project_required",
                message="At least one project row is required.",
            )
        )
        return
    for row_number, raw in rows:
        name = _cell_text(raw.get("Project name *"))
        developer = _cell_text(raw.get("Developer name *"))
        address = _cell_text(raw.get("Address line 1 *"))
        pin_code = _cell_pin(raw.get("PIN code *"))
        city = _cell_text(raw.get("City *"))
        state = _cell_text(raw.get("State *"))
        if not all([name, developer, address, pin_code, city, state]):
            parsed.errors.append(
                WorkbookRowError(
                    sheet=SHEET_PROJECT,
                    row_number=row_number,
                    code="project_incomplete",
                    message="Project name, developer, address, PIN, city, and state are required.",
                    raw=raw,
                )
            )
            continue
        area_unit = (
            _lookup(AREA_UNIT_MAP, raw.get("Area is measured in *")) or MeasurementUnit.SQ_FT
        )
        property_types = _lookup(PROPERTY_TYPE_MAP, raw.get("Project type *")) or [
            PropertyType.RESIDENTIAL
        ]
        gstin = _cell_text(raw.get("GST number"))
        name_key = project_join_key(name)
        if any(item["name_key"] == name_key for item in parsed.projects):
            parsed.skipped.append(
                WorkbookRowError(
                    sheet=SHEET_PROJECT,
                    row_number=row_number,
                    code="project_duplicate_skipped",
                    message="Duplicate project name in the workbook was skipped.",
                    raw=raw,
                )
            )
            continue
        parsed.projects.append(
            {
                "sheet_row": row_number,
                "name": " ".join(name.split()),
                "name_key": name_key,
                "developer_name": developer,
                "gstin": gstin if len(gstin) == 15 else None,
                "address_line_1": address,
                "address_line_2": _cell_text(raw.get("Address line 2")) or None,
                "pin_code": pin_code,
                "city": city,
                "state": state,
                "country": _cell_text(raw.get("Country")) or "India",
                "property_types": property_types,
                "primary_measurement_unit": area_unit,
                "possession_date": _parse_possession(raw.get("Expected possession")),
                "units_count": _parse_int(raw.get("Approx. total units or plots"), 0) or 0,
                "latitude": _parse_float(raw.get("Latitude")),
                "longitude": _parse_float(raw.get("Longitude")),
            }
        )


def _parse_towers(workbook: Any, parsed: ParsedProjectWorkbook) -> None:
    """Parse the Towers sheet into create_tower payloads."""
    _, rows = _read_sheet(workbook, SHEET_TOWERS)
    for row_number, raw in rows:
        name = _cell_text(raw.get("Tower or block name *"))
        project_key = _resolve_project_key(parsed, raw.get("Project name *"))
        tower_type = _lookup(TOWER_TYPE_MAP, raw.get("Building use *"))
        if not name or not project_key or tower_type is None:
            parsed.errors.append(
                WorkbookRowError(
                    sheet=SHEET_TOWERS,
                    row_number=row_number,
                    code="tower_incomplete",
                    message="Tower name, project name, and building use are required.",
                    raw=raw,
                )
            )
            continue
        numbering = _lookup(NUMBERING_MAP, raw.get("How flats are numbered"))
        if numbering is None:
            numbering = UnitNumberingPattern.FLOOR_UNIT
        parsed.towers.append(
            {
                "sheet_row": row_number,
                "project_key": project_key,
                "name": name,
                "code": _slug(name, "tower"),
                "tower_type": tower_type,
                "basement_count": _parse_int(raw.get("Basement levels"), 0) or 0,
                "upper_floor_count": _parse_int(raw.get("Floors above ground *"), 0) or 0,
                "units_per_floor_default": _parse_int(raw.get("Usual homes per floor")),
                "has_wings": _cell_bool(raw.get("Has wings? *")),
                "numbering_pattern": numbering,
                "starting_unit_number": _parse_int(raw.get("First flat number"), 1) or 1,
                "custom_prefix": _cell_text(raw.get("Number prefix")) or None,
                "latitude": _parse_float(raw.get("Latitude")),
                "longitude": _parse_float(raw.get("Longitude")),
            }
        )


def _parse_wings(workbook: Any, parsed: ParsedProjectWorkbook) -> None:
    """Parse the Wings sheet, skipping empty or unmatched rows."""
    _, rows = _read_sheet(workbook, SHEET_WINGS)
    tower_keys = {
        (item["project_key"], item["name"].casefold())
        for item in parsed.towers
        if item.get("has_wings")
    }
    for row_number, raw in rows:
        wing_name = _cell_text(raw.get("Wing name *"))
        tower_name = _cell_text(raw.get("Tower or block name *"))
        project_key = _resolve_project_key(parsed, raw.get("Project name *"))
        if not wing_name or (project_key, tower_name.casefold()) not in tower_keys:
            parsed.skipped.append(
                WorkbookRowError(
                    sheet=SHEET_WINGS,
                    row_number=row_number,
                    code="wing_skipped",
                    message="Empty or unknown wing row was skipped.",
                    raw=raw,
                )
            )
            continue
        parsed.wings.append(
            {
                "sheet_row": row_number,
                "project_key": project_key,
                "tower_name": tower_name,
                "name": wing_name,
                "code": _slug(wing_name, "wing"),
            }
        )


def _parse_unit_types(workbook: Any, parsed: ParsedProjectWorkbook) -> None:
    """Parse the Unit types sheet into create_config payloads."""
    _, rows = _read_sheet(workbook, SHEET_UNIT_TYPES)
    for row_number, raw in rows:
        name = _cell_text(raw.get("Type name *"))
        kind = _lookup(KIND_MAP, raw.get("Kind *"))
        project_key = _resolve_project_key(parsed, raw.get("Project name *"))
        if not name or kind is None or not project_key:
            parsed.errors.append(
                WorkbookRowError(
                    sheet=SHEET_UNIT_TYPES,
                    row_number=row_number,
                    code="unit_type_incomplete",
                    message="Type name, kind, and project name are required.",
                    raw=raw,
                )
            )
            continue
        area = _parse_float(raw.get("Area *"))
        parsed.unit_types.append(
            {
                "sheet_row": row_number,
                "project_key": project_key,
                "name": name,
                "code": _slug(name, "config"),
                "config_kind": kind,
                "bedrooms": _parse_float(raw.get("Bedrooms")),
                "bathrooms": _parse_float(raw.get("Bathrooms")),
                "area_sqft": area,
                "carpet_area_sqft": _parse_float(raw.get("Carpet area")),
                "balconies": _parse_int(raw.get("Balconies"), 0) or 0,
                "default_facing": _lookup(FACING_MAP, raw.get("Usual facing")),
                "two_wheeler_parking_entitlement": _parse_int(
                    raw.get("Two-wheeler parking included"), 0
                )
                or 0,
                "four_wheeler_parking_entitlement": _parse_int(raw.get("Car parking included"), 0)
                or 0,
                "commercial_unit_type": _lookup(SHOP_TYPE_MAP, raw.get("Shop type")),
                "plot_type": _lookup(PLOT_TYPE_MAP, raw.get("Plot category")),
            }
        )


def _ensure_unit_type(
    parsed: ParsedProjectWorkbook,
    *,
    project_key: str,
    type_name: str,
    area_sqft: float | None,
    source_row: int,
) -> dict[str, Any]:
    """Reuse or auto-create a unit type referenced by the units sheet."""
    existing = next(
        (
            item
            for item in parsed.unit_types
            if item["project_key"] == project_key
            and item["name"].casefold() == type_name.casefold()
        ),
        None,
    )
    if existing:
        if existing.get("area_sqft") in (None, 0) and area_sqft:
            existing["area_sqft"] = area_sqft
        return existing
    created = {
        "sheet_row": source_row,
        "project_key": project_key,
        "name": type_name,
        "code": _slug(type_name, "config"),
        "config_kind": UnitConfigKind.APARTMENT,
        "bedrooms": _infer_bedrooms(type_name),
        "bathrooms": 2.0,
        "area_sqft": area_sqft or 1.0,
        "carpet_area_sqft": None,
        "balconies": 0,
        "default_facing": None,
        "two_wheeler_parking_entitlement": 0,
        "four_wheeler_parking_entitlement": 0,
        "commercial_unit_type": None,
        "plot_type": None,
        "auto_created": True,
    }
    parsed.unit_types.append(created)
    return created


def _parse_units(workbook: Any, parsed: ParsedProjectWorkbook) -> None:
    """Parse occupant sample rows; the full unit grid is generated later."""
    _, rows = _read_sheet(workbook, SHEET_UNITS)
    floors_seen: set[tuple[str, str, str, int]] = set()
    for row_number, raw in rows:
        project_key = _resolve_project_key(parsed, raw.get("Project name *"))
        tower_name = _cell_text(raw.get("Tower or block"))
        wing_name = _cell_text(raw.get("Wing")) or None
        unit_code = _cell_code(raw.get("Unit or plot number *"))
        type_name = _cell_text(raw.get("Unit type *"))
        parking_floor = _cell_bool(raw.get("Parking floor?"))
        floor_level = parse_floor_level(raw.get("Floor number"))
        if floor_level is None:
            floor_level = parse_floor_level(raw.get("Floor name"))
        is_parking_floor = parking_floor and not unit_code
        if is_parking_floor and tower_name and floor_level is not None:
            _add_floor(
                parsed,
                floors_seen,
                project_key=project_key,
                tower_name=tower_name,
                wing_name=wing_name,
                floor_level=floor_level,
                display_name=floor_display_name(floor_level),
                is_parking=True,
                sheet_row=row_number,
            )
            continue
        if not unit_code:
            parsed.skipped.append(
                WorkbookRowError(
                    sheet=SHEET_UNITS,
                    row_number=row_number,
                    code="unit_skipped",
                    message="Unit row without a unit/plot number was skipped.",
                    raw=raw,
                )
            )
            continue
        if not type_name:
            parsed.errors.append(
                WorkbookRowError(
                    sheet=SHEET_UNITS,
                    row_number=row_number,
                    code="unit_type_missing",
                    message="Unit type is required for a home or plot row.",
                    raw=raw,
                )
            )
            continue
        area = _parse_float(raw.get("Area"))
        _ensure_unit_type(
            parsed,
            project_key=project_key,
            type_name=type_name,
            area_sqft=area,
            source_row=row_number,
        )
        if tower_name and floor_level is not None:
            _add_floor(
                parsed,
                floors_seen,
                project_key=project_key,
                tower_name=tower_name,
                wing_name=wing_name,
                floor_level=floor_level,
                display_name=floor_display_name(floor_level),
                is_parking=False,
                sheet_row=row_number,
            )
        parsed.units.append(
            {
                "sheet_row": row_number,
                "project_key": project_key,
                "tower_name": tower_name or None,
                "wing_name": wing_name,
                "floor_level": floor_level,
                "code": unit_code,
                "unit_label": unit_code,
                "type_name": type_name,
                "status": _lookup(UNIT_STATUS_MAP, raw.get("Status *")) or UnitStatus.VACANT,
                "is_parking": False,
            }
        )


def _add_floor(
    parsed: ParsedProjectWorkbook,
    floors_seen: set[tuple[str, str, str, int]],
    *,
    project_key: str,
    tower_name: str,
    wing_name: str | None,
    floor_level: int,
    display_name: str,
    is_parking: bool,
    sheet_row: int,
) -> None:
    """Record a unique floor for a tower, marking parking floors when needed."""
    key = (project_key, tower_name.casefold(), (wing_name or "").casefold(), floor_level)
    if key in floors_seen:
        if is_parking:
            for item in parsed.floors:
                item_key = (
                    item["project_key"],
                    item["tower_name"].casefold(),
                    (item.get("wing_name") or "").casefold(),
                    item["level_number"],
                )
                if item_key == key:
                    item["is_parking"] = True
                    break
        return
    floors_seen.add(key)
    parsed.floors.append(
        {
            "sheet_row": sheet_row,
            "project_key": project_key,
            "tower_name": tower_name,
            "wing_name": wing_name,
            "level_number": floor_level,
            "display_name": display_name,
            "is_parking": is_parking,
        }
    )


def _parse_amenities(workbook: Any, parsed: ParsedProjectWorkbook) -> None:
    """Parse the Amenities sheet into create_facility payloads."""
    _, rows = _read_sheet(workbook, SHEET_AMENITIES)
    tower_names = {(item["project_key"], item["name"].casefold()) for item in parsed.towers}
    for row_number, raw in rows:
        name = _cell_text(raw.get("Amenity name *"))
        project_key = _resolve_project_key(parsed, raw.get("Project name *"))
        facility_type = _lookup(AMENITY_TYPE_MAP, raw.get("Category *")) or FacilityType.RECREATION
        location_type = (
            _lookup(AMENITY_PLACE_MAP, raw.get("Where it sits"))
            or FacilityLocationType.INDOOR_CLUBHOUSE
        )
        if not name or not project_key:
            parsed.errors.append(
                WorkbookRowError(
                    sheet=SHEET_AMENITIES,
                    row_number=row_number,
                    code="amenity_incomplete",
                    message="Amenity name and project name are required.",
                    raw=raw,
                )
            )
            continue
        tower_name = _cell_text(raw.get("Tower or block"))
        tower_ref = tower_name if (project_key, tower_name.casefold()) in tower_names else None
        if location_type == FacilityLocationType.IN_TOWER and not tower_ref:
            location_type = FacilityLocationType.INDOOR_CLUBHOUSE
        parsed.amenities.append(
            {
                "sheet_row": row_number,
                "project_key": project_key,
                "name": name,
                "facility_type": facility_type,
                "location_type": location_type,
                "tower_name": tower_ref,
                "wing": _cell_text(raw.get("Wing")) or None,
                "floor_level": parse_floor_level(raw.get("Floor number")),
                "capacity_persons": _parse_int(raw.get("How many people")),
                "parking_slots": _parse_int(raw.get("Parking spaces")),
                "parking_user_type": _lookup(PARKING_USERS_MAP, raw.get("Who can use the parking")),
                "parking_vehicle_category": _lookup(VEHICLE_MAP, raw.get("Vehicle")),
                "latitude": _parse_float(raw.get("Latitude")),
                "longitude": _parse_float(raw.get("Longitude")),
            }
        )


def _parse_parking_zones(workbook: Any, parsed: ParsedProjectWorkbook) -> None:
    """Parse parking zones, skipping incomplete slot ranges."""
    _, rows = _read_sheet(workbook, SHEET_PARKING)
    for row_number, raw in rows:
        zone_name = _cell_text(raw.get("Zone name *"))
        slot_from = _parse_int(raw.get("First slot number *"))
        slot_to = _parse_int(raw.get("Last slot number *"))
        if not zone_name or slot_from is None or slot_to is None:
            parsed.skipped.append(
                WorkbookRowError(
                    sheet=SHEET_PARKING,
                    row_number=row_number,
                    code="parking_zone_skipped",
                    message="Incomplete parking zone row was skipped.",
                    raw=raw,
                )
            )
            continue
        parsed.parking_zones.append(
            {
                "sheet_row": row_number,
                "project_key": _resolve_project_key(parsed, raw.get("Project name *")),
                "tower_name": _cell_text(raw.get("Tower or block")) or None,
                "name": zone_name,
                "floor_level": parse_floor_level(raw.get("Floor number")),
                "slot_from": slot_from,
                "slot_to": slot_to,
            }
        )


def _parse_gates_lifts(workbook: Any, parsed: ParsedProjectWorkbook) -> None:
    """Parse gate and lift rows from the combined sheet."""
    _, rows = _read_sheet(workbook, SHEET_GATES_LIFTS)
    for row_number, raw in rows:
        kind = _cell_text(raw.get("This row is a *")).casefold()
        name = _cell_text(raw.get("Name *"))
        project_key = _resolve_project_key(parsed, raw.get("Project name *"))
        if not name or kind not in {"gate", "lift"}:
            parsed.skipped.append(
                WorkbookRowError(
                    sheet=SHEET_GATES_LIFTS,
                    row_number=row_number,
                    code="gate_lift_skipped",
                    message="Incomplete gate/lift row was skipped.",
                    raw=raw,
                )
            )
            continue
        common = {
            "sheet_row": row_number,
            "project_key": project_key,
            "tower_name": _cell_text(raw.get("Tower or block")) or None,
            "wing_name": _cell_text(raw.get("Wing")) or None,
            "name": name,
        }
        if kind == "gate":
            parsed.gates.append(
                {
                    **common,
                    "gate_type": _lookup(GATE_TYPE_MAP, raw.get("Gate use")) or GateType.BOTH,
                    "status": _lookup(GATE_STATUS_MAP, raw.get("Gate today")) or GateStatus.ACTIVE,
                    "is_open_24x7": _cell_bool(raw.get("Open 24 hours?")),
                }
            )
            continue
        parsed.lifts.append(
            {
                **common,
                "lift_type": _lookup(LIFT_TYPE_MAP, raw.get("Lift type")) or LiftType.PASSENGER,
                "capacity_persons": _parse_int(raw.get("Persons per lift")),
            }
        )


def _skip_optional_sheet(
    workbook: Any, parsed: ParsedProjectWorkbook, sheet_name: str, code: str
) -> None:
    """Skip optional media/map sheets and record them as skipped rows."""
    _, rows = _read_sheet(workbook, sheet_name)
    for row_number, raw in rows:
        parsed.skipped.append(
            WorkbookRowError(
                sheet=sheet_name,
                row_number=row_number,
                code=code,
                message="Optional file/map rows are not applied during project import.",
                raw=raw,
            )
        )
