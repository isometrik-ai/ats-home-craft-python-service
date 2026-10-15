"""Unit tests for Project_data_bulk workbook parsing rules."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from apps.user_service.app.schemas.enums import (
    MeasurementUnit,
    UnitConfigKind,
    UnitNumberingPattern,
    UnitStatus,
)
from apps.user_service.app.services.project_workbook_parser import (
    HEADER_ROW,
    build_tower_floor_specs,
    floor_display_name,
    normalize_project_name,
    parse_project_workbook,
    project_join_key,
    unit_codes_for_floor,
)


def _write_sheet(
    workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]
) -> None:
    sheet = workbook.create_sheet(title)
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=HEADER_ROW, column=index, value=header)
    for row_offset, values in enumerate(rows):
        for column, value in enumerate(values, start=1):
            sheet.cell(row=HEADER_ROW + 2 + row_offset, column=column, value=value)


def _workbook_bytes(sheets: dict[str, tuple[list[str], list[list[object]]]]) -> bytes:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for title, (headers, rows) in sheets.items():
        _write_sheet(workbook, title, headers, rows)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_normalize_project_name_collapses_and_casefolds():
    """Project names join across sheets after whitespace and case folding."""
    assert normalize_project_name("  Happy   Trails ") == normalize_project_name("happy trails")


def test_project_join_key_ignores_internal_spaces():
    """Home Kraft and Homekraft resolve to the same import join key."""
    assert project_join_key("ATS Home Kraft Happy Trails") == project_join_key(
        "ATS Homekraft Happy Trails"
    )
    assert project_join_key("Happy Trails") in project_join_key("ATS Home Kraft Happy Trails")


def test_floor_display_name_uses_floor_number_not_words():
    """Import labels Ground, Floor N, and Basement N from the numeric level."""
    assert floor_display_name(0) == "Ground"
    assert floor_display_name(1) == "Floor 1"
    assert floor_display_name(22) == "Floor 22"
    assert floor_display_name(-1) == "Basement 1"
    assert floor_display_name(-2) == "Basement 2"


def test_unit_codes_for_floor_follow_tower_numbering():
    """Floor + unit uses 1011/1012 style when first flat encodes a 2-digit floor."""
    assert unit_codes_for_floor(
        starting_unit_number=1011,
        floor_level=1,
        units_per_floor=2,
        numbering_pattern=UnitNumberingPattern.FLOOR_UNIT,
    ) == ["1011", "1012"]
    assert unit_codes_for_floor(
        starting_unit_number=1011,
        floor_level=22,
        units_per_floor=2,
        numbering_pattern=UnitNumberingPattern.FLOOR_UNIT,
    ) == ["1221", "1222"]
    assert unit_codes_for_floor(
        starting_unit_number=11011,
        floor_level=1,
        units_per_floor=8,
        numbering_pattern=UnitNumberingPattern.FLOOR_UNIT,
    ) == [str(11011 + index) for index in range(8)]
    assert unit_codes_for_floor(
        starting_unit_number=101,
        floor_level=2,
        units_per_floor=2,
        numbering_pattern=UnitNumberingPattern.FLOOR_UNIT,
    ) == ["201", "202"]


def test_build_tower_floor_specs_matches_create_tower_nested_floors():
    """Import passes the same nested floor payload create_tower already consumes."""
    specs = build_tower_floor_specs(basement_count=1, upper_floor_count=27, parking_levels={-1})
    by_level = {item["level_number"]: item for item in specs}
    assert by_level[-1] == {
        "level_number": -1,
        "display_name": "Basement 1",
        "is_parking": True,
    }
    assert by_level[1]["display_name"] == "Floor 1"
    assert by_level[22]["display_name"] == "Floor 22"
    assert by_level[27]["display_name"] == "Floor 27"
    assert 0 not in by_level


def test_parser_applies_happy_trails_import_rules():
    """Parser normalizes names, defaults units, auto-creates types, and skips junk."""
    payload = _workbook_bytes(
        {
            "1. Project": (
                [
                    "Project name *",
                    "Developer name *",
                    "Address line 1 *",
                    "PIN code *",
                    "City *",
                    "State *",
                    "Country",
                    "Area is measured in *",
                    "Project type *",
                    "GST number",
                ],
                [
                    [
                        "  Happy   Trails",
                        "ATS",
                        "Sector 1",
                        203207.0,
                        "Noida",
                        "UP",
                        "India",
                        "17,10,",
                        "Residential apartments",
                        "",
                    ],
                    [
                        "Happy Trails",
                        "ATS",
                        "Sector 1",
                        203207.0,
                        "Noida",
                        "UP",
                        "India",
                        "sq_ft",
                        "Residential apartments",
                        "",
                    ],
                ],
            ),
            "2. Towers": (
                [
                    "Tower or block name *",
                    "Project name *",
                    "Building use *",
                    "Has wings? *",
                    "Basement levels",
                    "Floors above ground *",
                ],
                [["Tower A", "happy trails", "Residential", "No", 1, 12]],
            ),
            "3. Wings": (
                ["Wing name *", "Tower or block name *", "Project name *"],
                [["", "Tower -1", "happy trails"], ["Tower -1", "Missing Tower", "happy trails"]],
            ),
            "4. Unit types": (
                ["Type name *", "Kind *", "Project name *", "Area *", "Bedrooms", "Bathrooms"],
                [],
            ),
            "5. Units and plots": (
                [
                    "Project name *",
                    "Tower or block",
                    "Wing",
                    "Unit or plot number *",
                    "Unit type *",
                    "Parking floor?",
                    "Floor number",
                    "Floor name",
                    "Area",
                    "Status *",
                    "Name shown to buyers",
                ],
                [
                    [
                        "Happy Trails",
                        "Tower A",
                        "",
                        "A-101",
                        "3+Study",
                        "Yes",
                        1,
                        "First Floor",
                        1450,
                        "Owner",
                        "Ravi Kumar",
                    ],
                    [
                        "Happy Trails",
                        "Tower A",
                        "",
                        "",
                        "",
                        "Yes",
                        -1,
                        "Basement",
                        None,
                        "",
                        "",
                    ],
                ],
            ),
            "7. Parking zones": (
                [
                    "Zone name *",
                    "First slot number *",
                    "Last slot number *",
                    "Project name *",
                    "Tower or block",
                ],
                [["P1", None, None, "Happy Trails", "Tower A"]],
            ),
        }
    )

    parsed = parse_project_workbook(payload)

    assert len(parsed.projects) == 1
    assert any(item.code == "project_duplicate_skipped" for item in parsed.skipped)
    assert parsed.projects[0]["gstin"] is None
    assert parsed.projects[0]["primary_measurement_unit"] == MeasurementUnit.SQ_FT
    assert parsed.projects[0]["pin_code"] == "203207"
    assert parsed.projects[0]["name_key"] == project_join_key("Happy Trails")
    assert parsed.towers[0]["project_key"] == parsed.projects[0]["name_key"]
    assert parsed.wings == []
    assert any(item.code == "wing_skipped" for item in parsed.skipped)
    assert any(item.code == "parking_zone_skipped" for item in parsed.skipped)
    assert len(parsed.units) == 1
    assert parsed.units[0]["code"] == "A-101"
    assert parsed.units[0]["status"] == UnitStatus.VACANT
    assert "occupant" not in parsed.units[0]
    assert "Ravi" not in str(parsed.units[0])
    auto_type = next(item for item in parsed.unit_types if item["name"] == "3+Study")
    assert auto_type["auto_created"] is True
    assert auto_type["config_kind"] == UnitConfigKind.APARTMENT
    assert auto_type["bedrooms"] == 3.0
    parking_floors = [item for item in parsed.floors if item["is_parking"]]
    home_floors = [item for item in parsed.floors if not item["is_parking"]]
    assert parking_floors[0]["level_number"] == -1
    assert parking_floors[0]["display_name"] == "Basement 1"
    assert parsed.towers[0]["basement_count"] == 1
    assert parsed.towers[0]["upper_floor_count"] == 12
    assert home_floors[0]["level_number"] == 1
    assert home_floors[0]["display_name"] == "Floor 1"
    assert home_floors[0]["is_parking"] is False
    assert "First Floor" not in {item["display_name"] for item in parsed.floors}


def test_parser_joins_homecraft_spelling_variants():
    """Towers using Homekraft still attach to a Home Kraft project row."""
    payload = _workbook_bytes(
        {
            "1. Project": (
                [
                    "Project name *",
                    "Developer name *",
                    "Address line 1 *",
                    "PIN code *",
                    "City *",
                    "State *",
                ],
                [["ATS Home Kraft Happy Trails", "ATS", "Sector 1", "203207", "Noida", "UP"]],
            ),
            "2. Towers": (
                [
                    "Tower or block name *",
                    "Project name *",
                    "Building use *",
                    "Has wings? *",
                    "Basement levels",
                    "Floors above ground *",
                ],
                [["Tower 1", "ATS Homekraft Happy Trails", "Residential", "No", 1, 27]],
            ),
            "4. Unit types": (
                ["Type name *", "Kind *", "Project name *", "Area *", "Bedrooms", "Bathrooms"],
                [["2 BHK", "Apartment", "Happy Trails", 1100, 2, 2]],
            ),
            "5. Units and plots": (
                [
                    "Project name *",
                    "Tower or block",
                    "Unit or plot number *",
                    "Unit type *",
                    "Floor number",
                    "Floor name",
                    "Status *",
                ],
                [
                    [
                        "ATS Homekraft Happy Trails",
                        "Tower 1",
                        "1221",
                        "3+Study",
                        22,
                        "Twenty two",
                        "Vacant",
                    ]
                ],
            ),
        }
    )
    parsed = parse_project_workbook(payload)
    key = parsed.projects[0]["name_key"]
    assert parsed.towers[0]["project_key"] == key
    assert parsed.units[0]["project_key"] == key
    assert parsed.unit_types[0]["project_key"] == key
    assert any(item["name"] == "3+Study" for item in parsed.unit_types)
    floor_22 = next(item for item in parsed.floors if item["level_number"] == 22)
    assert floor_22["display_name"] == "Floor 22"
    assert not any(item["display_name"] == "Twenty two" for item in parsed.floors)


def test_parser_strips_excel_float_unit_codes():
    """Excel numeric unit cells import as 1221, not 1221.0."""
    payload = _workbook_bytes(
        {
            "1. Project": (
                [
                    "Project name *",
                    "Developer name *",
                    "Address line 1 *",
                    "PIN code *",
                    "City *",
                    "State *",
                ],
                [["ATS Home Kraft Happy Trails", "ATS", "Sector 1", "203207", "Noida", "UP"]],
            ),
            "2. Towers": (
                [
                    "Project name *",
                    "Tower or block name *",
                    "Building use *",
                    "Has wings? *",
                    "Floors above ground *",
                ],
                [["Tower 1", "ATS Homekraft Happy Trails", "Residential", "No", 27]],
            ),
            "5. Units and plots": (
                [
                    "Project name *",
                    "Tower or block",
                    "Unit or plot number *",
                    "Unit type *",
                    "Floor number",
                    "Status *",
                ],
                [["ATS Homekraft Happy Trails", "Tower 1", 1221.0, "3+Study", 22, "Owner"]],
            ),
        }
    )
    parsed = parse_project_workbook(payload)
    assert parsed.units[0]["code"] == "1221"
